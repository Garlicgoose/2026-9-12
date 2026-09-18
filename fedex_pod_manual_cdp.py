# -*- coding: utf-8 -*-
"""FedEx POD 半自动归档（推荐版：CDP / Page.printToPDF）

使用方式：
1. 用户自己启动 Edge 或 Chrome（需开启 remote debugging），并自己打开 FedEx 查询网页。
2. 本程序只连接已经打开的浏览器，不会自动导航、刷新、输入或提交 FedEx 查询。
3. 程序读取 Excel 运单号，把当前运单复制到剪贴板。
4. 用户在 FedEx 页面粘贴运单号并查询。
5. 程序检测到查询结果页后，自动把“当前已加载页面”打印为：运单号.pdf
6. 用户手动点击 FedEx 详情页/更多详情。
7. 程序检测到详情内容后，再自动打印为：运单号+.pdf
8. 完成后自动复制下一票运单号。
9. 页面异常、验证码、FedEx 服务错误、浏览器失联、PDF 保存失败时立即停止推进，
   不刷新、不重试 FedEx 请求；由用户在置顶面板点击“重试当前”或“跳过当前”。

依赖：
    py -m pip install openpyxl playwright

注意：
    Playwright 这里只用于连接你自己已经打开的 Chromium 浏览器以及读取当前页面/打印 PDF，
    不会由程序打开 FedEx 网站或提交查询。
"""
from __future__ import annotations

import argparse
import base64
import json
import os
import queue
import re
import socket
import sys
import threading
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Optional
from urllib.error import URLError
from urllib.request import urlopen

import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from tkinter.scrolledtext import ScrolledText


APP_TITLE = "FedEx POD 半自动助手 - CDP打印版"
TRACKING_RE = re.compile(r"^[A-Za-z0-9]{8,30}$")

TRACKING_HEADERS = {
    "运单号", "快递单号", "追踪号", "tracking number", "tracking no",
    "tracking no.", "awb", "awb no", "waybill", "fedex tracking number",
}

DETAIL_BUTTON_RE = re.compile(
    r"(?:view|see)\s+(?:more\s+|full\s+)?(?:shipment\s+)?details?|"
    r"shipment\s+details?|detailed\s+results?|travel\s+history|"
    r"查看(?:更多)?详细信息|查看(?:更多)?详情|货件详情|详细信息",
    re.I,
)

DETAIL_CONTENT_RE = re.compile(
    r"travel\s+history|shipment\s+facts|shipment\s+details|"
    r"运输历史记录|物流记录|行程历史|货件详情|货件事实|详细信息",
    re.I,
)

BLOCK_PAGE_RE = re.compile(
    r"too many requests|access denied|temporarily unavailable|service unavailable|"
    r"unusual traffic|try again later|captcha|robot|429|503|"
    r"请求过多|访问被拒绝|暂时不可用|稍后重试|验证码|系统(?:错误|繁忙)",
    re.I,
)

EDGE_PATHS = (
    Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
    Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
)
CHROME_PATHS = (
    Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe"),
    Path(r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe"),
    Path(os.environ.get("LOCALAPPDATA", "")) / "Google/Chrome/Application/chrome.exe",
)

DEFAULT_PORTS = {"edge": 9222, "chrome": 9223}


class FedExManualError(RuntimeError):
    pass


@dataclass
class Job:
    row: int
    number: str


@dataclass
class WorkbookState:
    source: Path
    result_path: Path
    output_dir: Path
    wb: Any
    ws: Any
    result_cols: dict[str, int]
    jobs: list[Job]


def normalize_tracking_number(value: Any) -> str:
    if value is None:
        raise ValueError("空运单号")
    if isinstance(value, bool):
        raise ValueError(f"FedEx 运单号格式不正确：{value!r}")
    if isinstance(value, int):
        raw = str(value)
    elif isinstance(value, float) and value.is_integer():
        raw = str(int(value))
    else:
        raw = str(value).strip()
        if raw.endswith(".0") and raw[:-2].isdigit():
            raw = raw[:-2]
    number = re.sub(r"[\s-]+", "", raw)
    if not TRACKING_RE.fullmatch(number):
        raise ValueError(f"FedEx 运单号格式不正确：{value!r}")
    return number


def output_paths(output_dir: Path, tracking_number: str) -> tuple[Path, Path]:
    number = normalize_tracking_number(tracking_number)
    return output_dir / f"{number}.pdf", output_dir / f"{number}+.pdf"


def is_valid_pdf(path: Path) -> bool:
    try:
        return (
            path.is_file()
            and path.stat().st_size >= 1024
            and path.read_bytes()[:5] == b"%PDF-"
        )
    except OSError:
        return False


def clean_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def find_tracking_column(ws, requested: str = "") -> tuple[int, int]:
    wanted = clean_header(requested)
    for row in range(1, min(ws.max_row, 20) + 1):
        for col in range(1, ws.max_column + 1):
            header = clean_header(ws.cell(row, col).value)
            if wanted and header == wanted:
                return row, col
            if not wanted and header in TRACKING_HEADERS:
                return row, col
    if requested:
        raise ValueError(f"找不到指定运单号列：{requested}")
    raise ValueError(
        "找不到运单号列。支持：运单号、快递单号、追踪号、Tracking Number、AWB 等"
    )


def ensure_result_columns(ws, header_row: int) -> dict[str, int]:
    names = ["POD状态", "POD主页文件", "POD详情文件", "POD错误", "POD处理时间"]
    found: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        name = str(ws.cell(header_row, col).value or "").strip()
        if name in names:
            found[name] = col
    for name in names:
        if name not in found:
            col = ws.max_column + 1
            ws.cell(header_row, col, name)
            found[name] = col
    return found


def load_workbook_state(
    excel_path: str,
    output_dir: str = "",
    sheet_name: str = "",
    tracking_column: str = "",
) -> WorkbookState:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("缺少 openpyxl，请执行：py -m pip install openpyxl") from exc

    source = Path(excel_path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(f"Excel 文件不存在：{source}")
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("仅支持 .xlsx / .xlsm")

    pod_dir = Path(output_dir).expanduser().resolve() if output_dir else source.parent / "FedEx_POD"
    pod_dir.mkdir(parents=True, exist_ok=True)
    result_path = source.with_name(f"{source.stem}_POD结果{source.suffix}")

    keep_vba = source.suffix.lower() == ".xlsm"
    wb = load_workbook(source, keep_vba=keep_vba)
    ws = wb[sheet_name] if sheet_name else wb.active
    header_row, tracking_col = find_tracking_column(ws, tracking_column)
    result_cols = ensure_result_columns(ws, header_row)

    jobs: list[Job] = []
    seen: set[str] = set()
    for row in range(header_row + 1, ws.max_row + 1):
        raw = ws.cell(row, tracking_col).value
        if raw is None or str(raw).strip() == "":
            continue
        try:
            number = normalize_tracking_number(raw)
        except ValueError as exc:
            ws.cell(row, result_cols["POD状态"], "跳过")
            ws.cell(row, result_cols["POD错误"], str(exc))
            continue
        if number in seen:
            ws.cell(row, result_cols["POD状态"], "跳过重复")
            continue
        seen.add(number)

        main_pdf, detail_pdf = output_paths(pod_dir, number)
        if is_valid_pdf(main_pdf) and is_valid_pdf(detail_pdf):
            ws.cell(row, result_cols["POD状态"], "已存在")
            ws.cell(row, result_cols["POD主页文件"], str(main_pdf))
            ws.cell(row, result_cols["POD详情文件"], str(detail_pdf))
            continue
        jobs.append(Job(row=row, number=number))

    wb.save(result_path)
    return WorkbookState(
        source=source,
        result_path=result_path,
        output_dir=pod_dir,
        wb=wb,
        ws=ws,
        result_cols=result_cols,
        jobs=jobs,
    )


def find_browser_path(browser: str) -> str:
    browser = browser.lower()
    candidates = CHROME_PATHS if browser == "chrome" else EDGE_PATHS
    for p in candidates:
        if p and p.is_file():
            return str(p.resolve())
    return "chrome.exe" if browser == "chrome" else "msedge.exe"


def build_debug_command(browser: str, port: int) -> str:
    exe = find_browser_path(browser)
    base = Path(os.environ.get("LOCALAPPDATA") or str(Path.home())) / "FedExManualBrowser"
    profile = base / ("Chrome" if browser == "chrome" else "Edge")
    return f'"{exe}" --remote-debugging-port={port} --user-data-dir="{profile}"'


def wait_for_cdp(cdp_url: str, timeout_seconds: float = 4.0) -> None:
    deadline = time.monotonic() + timeout_seconds
    last_error: Optional[BaseException] = None
    while time.monotonic() < deadline:
        try:
            with urlopen(f"{cdp_url}/json/version", timeout=1) as response:
                data = json.loads(response.read().decode("utf-8"))
            if data.get("webSocketDebuggerUrl"):
                return
        except (OSError, ValueError, URLError) as exc:
            last_error = exc
        time.sleep(0.25)
    raise FedExManualError(f"无法连接浏览器调试端口 {cdp_url}：{last_error}")


def body_text(page) -> str:
    try:
        return page.locator("body").inner_text(timeout=2500)
    except Exception:
        return ""


def normalized_text(text: str) -> str:
    return re.sub(r"[\s-]+", "", text or "")


def visible_detail_candidate(page, tracking_number: str = ""):
    locators = []
    for role in ("button", "link"):
        try:
            locators.append(page.get_by_role(role, name=DETAIL_BUTTON_RE))
        except Exception:
            pass
    try:
        locators.append(page.get_by_text(DETAIL_BUTTON_RE))
    except Exception:
        pass

    fallback = None
    for locator in locators:
        try:
            count = min(locator.count(), 20)
        except Exception:
            continue
        for i in range(count):
            try:
                candidate = locator.nth(i)
                if not candidate.is_visible():
                    continue
                if fallback is None:
                    fallback = candidate
                if tracking_number and candidate.evaluate(
                    """(element, number) => {
                        const wanted = String(number).replace(/[\\s-]+/g, '');
                        let current = element;
                        for (let depth = 0; current && depth < 8; depth += 1) {
                            const text = (current.innerText || '').replace(/[\\s-]+/g, '');
                            if (text.includes(wanted)) return true;
                            current = current.parentElement;
                        }
                        return false;
                    }""",
                    tracking_number,
                ):
                    return candidate
            except Exception:
                continue
    return fallback


def main_page_ready(page, tracking_number: str, text: str) -> bool:
    if tracking_number not in normalized_text(text):
        return False
    # 以可见“详情”入口作为“查询结果已经稳定显示”的强信号。
    return visible_detail_candidate(page, tracking_number) is not None


def detail_page_ready(page, tracking_number: str, text: str, main_text: str, main_url: str) -> bool:
    if tracking_number not in normalized_text(text):
        return False
    if text == main_text:
        return False
    changed_url = str(page.url) != str(main_url)
    content_signal = bool(DETAIL_CONTENT_RE.search(text))
    # 详情页可能在当前页展开，也可能跳转到另一 URL。
    return content_signal and (changed_url or abs(len(text) - len(main_text)) >= 80)


def dismiss_cookie_banner(page) -> None:
    pattern = re.compile(
        r"reject optional cookies|accept all cookies|拒绝可选|仅使用必要|接受全部",
        re.I,
    )
    try:
        locator = page.get_by_role("button", name=pattern)
        for i in range(min(locator.count(), 6)):
            candidate = locator.nth(i)
            if candidate.is_visible():
                candidate.click(timeout=1200)
                page.wait_for_timeout(200)
                return
    except Exception:
        pass


def hide_print_overlays(page) -> None:
    dismiss_cookie_banner(page)
    try:
        page.add_style_tag(content="""
            #usercentrics-root,
            iframe[src*="usercentrics"],
            iframe[src*="nuance"],
            [id*="nuance"],
            [class*="nuance"] {
                display: none !important;
                visibility: hidden !important;
            }
        """)
        page.wait_for_timeout(150)
    except Exception:
        pass


def print_current_page(context, page, destination: Path) -> None:
    if "fedex.com" not in str(page.url).casefold():
        raise FedExManualError("当前不是 FedEx 官网页面")
    destination.parent.mkdir(parents=True, exist_ok=True)
    hide_print_overlays(page)
    temporary = destination.with_name(f".{destination.name}.{os.getpid()}.tmp")
    session = context.new_cdp_session(page)
    try:
        result = session.send(
            "Page.printToPDF",
            {
                "landscape": False,
                "displayHeaderFooter": False,
                "printBackground": True,
                "preferCSSPageSize": False,
                "paperWidth": 8.27,
                "paperHeight": 11.69,
                "marginTop": 0.25,
                "marginBottom": 0.25,
                "marginLeft": 0.25,
                "marginRight": 0.25,
            },
        )
        temporary.write_bytes(base64.b64decode(result["data"], validate=True))
        if not is_valid_pdf(temporary):
            raise FedExManualError(f"生成的 PDF 无效：{destination.name}")
        temporary.replace(destination)
    finally:
        try:
            session.detach()
        except Exception:
            pass
        temporary.unlink(missing_ok=True)


class ManualWorker(threading.Thread):
    def __init__(
        self,
        ui_queue: queue.Queue,
        excel_path: str,
        output_dir: str,
        browser_name: str,
        port: int,
        sheet_name: str = "",
        tracking_column: str = "",
    ) -> None:
        super().__init__(daemon=True)
        self.ui_queue = ui_queue
        self.excel_path = excel_path
        self.output_dir = output_dir
        self.browser_name = browser_name
        self.port = port
        self.sheet_name = sheet_name
        self.tracking_column = tracking_column
        self.pause_event = threading.Event()
        self.stop_event = threading.Event()
        self.retry_event = threading.Event()
        self.skip_event = threading.Event()
        self.state: Optional[WorkbookState] = None
        self.playwright = None
        self.browser = None

    def post(self, kind: str, **data: Any) -> None:
        self.ui_queue.put((kind, data))

    def save_row(
        self,
        job: Job,
        status: str,
        main_pdf: str = "",
        detail_pdf: str = "",
        error: str = "",
    ) -> None:
        assert self.state is not None
        ws = self.state.ws
        cols = self.state.result_cols
        ws.cell(job.row, cols["POD状态"], status)
        if main_pdf:
            ws.cell(job.row, cols["POD主页文件"], main_pdf)
        if detail_pdf:
            ws.cell(job.row, cols["POD详情文件"], detail_pdf)
        ws.cell(job.row, cols["POD错误"], error)
        ws.cell(job.row, cols["POD处理时间"], time.strftime("%Y-%m-%d %H:%M:%S"))
        self.state.wb.save(self.state.result_path)

    def wait_if_paused(self) -> bool:
        while self.pause_event.is_set() and not self.stop_event.is_set():
            time.sleep(0.15)
        return not self.stop_event.is_set()

    def fault(self, message: str) -> str:
        self.post("error", message=message)
        self.retry_event.clear()
        self.skip_event.clear()
        while not self.stop_event.is_set():
            if self.skip_event.wait(0.15):
                self.skip_event.clear()
                return "skip"
            if self.retry_event.is_set():
                self.retry_event.clear()
                return "retry"
        return "stop"

    def get_fedex_pages(self):
        if self.browser is None:
            return []
        pages = []
        for context in self.browser.contexts:
            for page in context.pages:
                try:
                    if "fedex.com" in str(page.url).casefold():
                        pages.append((context, page))
                except Exception:
                    continue
        return pages

    def find_page_for_number(self, number: str):
        pages = self.get_fedex_pages()
        fallback = pages[-1] if pages else None
        for context, page in reversed(pages):
            text = body_text(page)
            if number in normalized_text(text):
                return context, page, text
        if fallback:
            context, page = fallback
            return context, page, body_text(page)
        return None

    def process_one(self, job: Job, index: int, total: int) -> str:
        assert self.state is not None
        number = job.number
        main_pdf, detail_pdf = output_paths(self.state.output_dir, number)
        main_done = is_valid_pdf(main_pdf)
        detail_done = is_valid_pdf(detail_pdf)

        self.post("current", number=number, index=index, total=total)
        self.post("clipboard", text=number)
        self.post(
            "status",
            message=(
                "主页已存在：请直接查询该单号并打开详情页" if main_done and not detail_done
                else "已复制运单号：请在 FedEx 粘贴并查询"
            ),
        )

        if detail_done and main_done:
            self.save_row(job, "已存在", str(main_pdf), str(detail_pdf))
            return "done"

        stage = "detail" if main_done else "main"
        main_text = ""
        main_url = ""
        mismatch_since: Optional[float] = None

        while not self.stop_event.is_set():
            if not self.wait_if_paused():
                return "stop"

            try:
                found = self.find_page_for_number(number)
            except Exception as exc:
                action = self.fault(f"浏览器连接异常：{exc}")
                if action == "skip":
                    self.save_row(job, "跳过", str(main_pdf) if main_done else "", "", str(exc))
                    return "skip"
                if action == "stop":
                    return "stop"
                continue

            if found is None:
                self.post("status", message="等待你打开 FedEx 查询页面……")
                time.sleep(0.5)
                continue

            context, page, text = found

            if BLOCK_PAGE_RE.search(text):
                action = self.fault("检测到 FedEx 限流/验证码/服务异常。已停止推进，不会自动刷新或重试。")
                if action == "skip":
                    self.save_row(job, "跳过", str(main_pdf) if main_done else "", "", "FedEx 页面异常")
                    return "skip"
                if action == "stop":
                    return "stop"
                continue

            if stage == "main":
                if number not in normalized_text(text):
                    self.post("status", message="等待你粘贴当前运单并查询……")
                    time.sleep(0.45)
                    continue
                if not main_page_ready(page, number, text):
                    self.post("status", message="已看到当前运单，等待查询结果稳定显示……")
                    time.sleep(0.45)
                    continue
                try:
                    self.post("status", message="检测到查询结果，正在保存主页 PDF……")
                    main_text = text
                    main_url = str(page.url)
                    print_current_page(context, page, main_pdf)
                    main_done = True
                    self.save_row(job, "主页已保存，等待详情", str(main_pdf), "", "")
                    self.post("status", message="✓ 主页已保存。请手动点击 FedEx 详情页/更多详情。")
                    stage = "detail"
                    mismatch_since = None
                except Exception as exc:
                    action = self.fault(f"主页 PDF 保存失败：{exc}")
                    if action == "skip":
                        self.save_row(job, "跳过", str(main_pdf) if is_valid_pdf(main_pdf) else "", "", str(exc))
                        return "skip"
                    if action == "stop":
                        return "stop"
                time.sleep(0.5)
                continue

            # stage == detail
            if not main_text:
                # 断点续跑时主页 PDF 已有，但程序没有本次页面快照。
                if number in normalized_text(text) and main_page_ready(page, number, text):
                    main_text = text
                    main_url = str(page.url)
                    self.post("status", message="主页 PDF 已存在。已识别查询页，请点击详情。")
                    time.sleep(0.5)
                    continue
                # 若用户已经直接进入详情，也允许直接识别。
                if number in normalized_text(text) and DETAIL_CONTENT_RE.search(text):
                    main_text = "__existing_main_pdf__"
                    main_url = ""

            if number not in normalized_text(text):
                if mismatch_since is None:
                    mismatch_since = time.monotonic()
                elif time.monotonic() - mismatch_since > 6:
                    action = self.fault(
                        f"当前 FedEx 页面连续 6 秒未检测到运单 {number}。已停止，避免把别的页面保存到当前单号。"
                    )
                    mismatch_since = None
                    if action == "skip":
                        self.save_row(job, "跳过", str(main_pdf) if main_done else "", "", "页面运单不匹配")
                        return "skip"
                    if action == "stop":
                        return "stop"
                time.sleep(0.4)
                continue
            mismatch_since = None

            ready = False
            if main_text == "__existing_main_pdf__":
                ready = bool(DETAIL_CONTENT_RE.search(text))
            elif main_text:
                ready = detail_page_ready(page, number, text, main_text, main_url)

            if not ready:
                self.post("status", message="✓ 主页已保存；等待你点击详情页……")
                time.sleep(0.45)
                continue

            try:
                self.post("status", message="检测到详情内容，正在保存详情 PDF……")
                print_current_page(context, page, detail_pdf)
                detail_done = True
                self.save_row(job, "成功", str(main_pdf), str(detail_pdf), "")
                self.post("status", message="✓ 当前运单两份 PDF 已保存，准备下一票。")
                time.sleep(0.6)
                return "done"
            except Exception as exc:
                action = self.fault(f"详情 PDF 保存失败：{exc}")
                if action == "skip":
                    self.save_row(job, "跳过", str(main_pdf), str(detail_pdf) if is_valid_pdf(detail_pdf) else "", str(exc))
                    return "skip"
                if action == "stop":
                    return "stop"

        return "stop"

    def run(self) -> None:
        try:
            self.post("status", message="正在读取 Excel……")
            self.state = load_workbook_state(
                self.excel_path,
                self.output_dir,
                self.sheet_name,
                self.tracking_column,
            )
            total = len(self.state.jobs)
            self.post(
                "loaded",
                total=total,
                result_path=str(self.state.result_path),
                output_dir=str(self.state.output_dir),
            )
            if total == 0:
                self.post("finished", message="没有需要处理的新运单。")
                return

            try:
                from playwright.sync_api import sync_playwright
            except ImportError as exc:
                raise RuntimeError("缺少 Playwright，请执行：py -m pip install playwright") from exc

            cdp_url = f"http://127.0.0.1:{self.port}"
            self.post("status", message=f"正在连接 {self.browser_name.title()}：{cdp_url}")
            wait_for_cdp(cdp_url)
            self.playwright = sync_playwright().start()
            self.browser = self.playwright.chromium.connect_over_cdp(cdp_url)
            self.post("status", message="浏览器已连接；不会自动打开、刷新或提交 FedEx 页面。")

            for idx, job in enumerate(self.state.jobs, 1):
                if self.stop_event.is_set():
                    break
                result = self.process_one(job, idx, total)
                if result == "stop":
                    break

            if self.stop_event.is_set():
                self.post("finished", message="已停止。已完成的 Excel/PDF 进度均已保存。")
            else:
                self.post("finished", message="本次队列处理结束。")

        except Exception as exc:
            self.post("fatal", message=str(exc))
        finally:
            try:
                if self.state is not None:
                    self.state.wb.save(self.state.result_path)
            except Exception:
                pass
            try:
                if self.browser is not None:
                    # connect_over_cdp 连接的是用户浏览器；不要关闭用户浏览器。
                    # 仅断开 Playwright 连接。
                    self.browser = None
            except Exception:
                pass
            try:
                if self.playwright is not None:
                    self.playwright.stop()
            except Exception:
                pass


class App:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title(APP_TITLE)
        self.root.geometry("390x455")
        self.root.minsize(360, 420)
        self.root.attributes("-topmost", True)
        self.root.protocol("WM_DELETE_WINDOW", self.on_close)

        self.ui_queue: queue.Queue = queue.Queue()
        self.worker: Optional[ManualWorker] = None

        self.excel_var = tk.StringVar()
        self.output_var = tk.StringVar()
        self.browser_var = tk.StringVar(value="Edge")
        self.port_var = tk.StringVar(value=str(DEFAULT_PORTS["edge"]))
        self.current_var = tk.StringVar(value="—")
        self.progress_var = tk.StringVar(value="0 / 0")
        self.status_var = tk.StringVar(value="选择 Excel 后点击开始")
        self.result_var = tk.StringVar(value="")

        self._build_ui()
        self.root.after(120, self.poll_ui_queue)

    def _build_ui(self) -> None:
        pad = {"padx": 8, "pady": 4}
        top = ttk.Frame(self.root)
        top.pack(fill="x", padx=8, pady=(8, 2))

        ttk.Label(top, text="Excel").grid(row=0, column=0, sticky="w")
        ttk.Entry(top, textvariable=self.excel_var).grid(row=0, column=1, sticky="ew", padx=5)
        ttk.Button(top, text="选择", command=self.pick_excel, width=7).grid(row=0, column=2)

        ttk.Label(top, text="PDF目录").grid(row=1, column=0, sticky="w")
        ttk.Entry(top, textvariable=self.output_var).grid(row=1, column=1, sticky="ew", padx=5)
        ttk.Button(top, text="选择", command=self.pick_output, width=7).grid(row=1, column=2)

        ttk.Label(top, text="浏览器").grid(row=2, column=0, sticky="w")
        browser_box = ttk.Combobox(
            top,
            textvariable=self.browser_var,
            values=["Edge", "Chrome"],
            state="readonly",
            width=10,
        )
        browser_box.grid(row=2, column=1, sticky="w", padx=5)
        browser_box.bind("<<ComboboxSelected>>", self.on_browser_changed)

        ttk.Label(top, text="调试端口").grid(row=3, column=0, sticky="w")
        ttk.Entry(top, textvariable=self.port_var, width=10).grid(row=3, column=1, sticky="w", padx=5)
        ttk.Button(top, text="复制启动命令", command=self.copy_debug_command).grid(row=3, column=2)
        top.columnconfigure(1, weight=1)

        ttk.Separator(self.root).pack(fill="x", padx=8, pady=5)

        current = ttk.Frame(self.root)
        current.pack(fill="x", padx=10)
        ttk.Label(current, text="当前运单", font=("Microsoft YaHei UI", 9)).pack(anchor="w")
        ttk.Label(
            current,
            textvariable=self.current_var,
            font=("Consolas", 18, "bold"),
        ).pack(anchor="center", pady=(2, 0))
        ttk.Label(current, textvariable=self.progress_var, font=("Microsoft YaHei UI", 11)).pack(anchor="center")
        self.progress = ttk.Progressbar(current, mode="determinate", maximum=100)
        self.progress.pack(fill="x", pady=(3, 5))

        status_frame = ttk.LabelFrame(self.root, text="状态")
        status_frame.pack(fill="x", padx=8, pady=4)
        ttk.Label(
            status_frame,
            textvariable=self.status_var,
            wraplength=350,
            justify="left",
        ).pack(fill="x", padx=8, pady=7)

        buttons = ttk.Frame(self.root)
        buttons.pack(fill="x", padx=8, pady=4)
        self.start_btn = ttk.Button(buttons, text="开始 / 连接", command=self.start)
        self.start_btn.grid(row=0, column=0, padx=2, sticky="ew")
        self.pause_btn = ttk.Button(buttons, text="暂停", command=self.toggle_pause, state="disabled")
        self.pause_btn.grid(row=0, column=1, padx=2, sticky="ew")
        self.retry_btn = ttk.Button(buttons, text="重试当前", command=self.retry_current, state="disabled")
        self.retry_btn.grid(row=0, column=2, padx=2, sticky="ew")
        self.skip_btn = ttk.Button(buttons, text="跳过当前", command=self.skip_current, state="disabled")
        self.skip_btn.grid(row=0, column=3, padx=2, sticky="ew")
        for i in range(4):
            buttons.columnconfigure(i, weight=1)

        lower = ttk.Frame(self.root)
        lower.pack(fill="x", padx=8, pady=(0, 4))
        ttk.Button(lower, text="重新复制当前单号", command=self.recopy_current).pack(side="left")
        ttk.Button(lower, text="停止", command=self.stop, width=9).pack(side="right")

        self.log = ScrolledText(self.root, height=6, font=("Consolas", 8), state="disabled")
        self.log.pack(fill="both", expand=True, padx=8, pady=(2, 8))

    def log_line(self, text: str) -> None:
        self.log.configure(state="normal")
        self.log.insert("end", time.strftime("%H:%M:%S") + "  " + text + "\n")
        self.log.see("end")
        self.log.configure(state="disabled")

    def pick_excel(self) -> None:
        path = filedialog.askopenfilename(
            title="选择包含 FedEx 运单号的 Excel",
            filetypes=[("Excel", "*.xlsx *.xlsm")],
        )
        if path:
            self.excel_var.set(path)
            if not self.output_var.get().strip():
                self.output_var.set(str(Path(path).resolve().parent / "FedEx_POD"))

    def pick_output(self) -> None:
        path = filedialog.askdirectory(title="选择 PDF 输出目录")
        if path:
            self.output_var.set(path)

    def on_browser_changed(self, _event=None) -> None:
        browser = self.browser_var.get().strip().lower()
        self.port_var.set(str(DEFAULT_PORTS.get(browser, 9222)))

    def copy_to_clipboard(self, text: str) -> None:
        try:
            self.root.clipboard_clear()
            self.root.clipboard_append(text)
            self.root.update_idletasks()
        except tk.TclError:
            pass

    def copy_debug_command(self) -> None:
        browser = self.browser_var.get().strip().lower()
        try:
            port = int(self.port_var.get().strip())
        except ValueError:
            messagebox.showerror(APP_TITLE, "调试端口必须是数字")
            return
        cmd = build_debug_command(browser, port)
        self.copy_to_clipboard(cmd)
        self.status_var.set("浏览器调试启动命令已复制。请关闭该浏览器后自行用此命令启动，再自行打开 FedEx。")
        self.log_line("已复制浏览器调试启动命令（程序不会自动启动浏览器）")

    def start(self) -> None:
        if self.worker and self.worker.is_alive():
            messagebox.showinfo(APP_TITLE, "任务已经在运行。")
            return
        excel = self.excel_var.get().strip()
        if not excel:
            self.pick_excel()
            excel = self.excel_var.get().strip()
        if not excel:
            return
        browser = self.browser_var.get().strip().lower()
        try:
            port = int(self.port_var.get().strip())
            if not (1 <= port <= 65535):
                raise ValueError
        except ValueError:
            messagebox.showerror(APP_TITLE, "调试端口必须是 1-65535 的数字")
            return

        self.worker = ManualWorker(
            self.ui_queue,
            excel,
            self.output_var.get().strip(),
            browser,
            port,
        )
        self.start_btn.configure(state="disabled")
        self.pause_btn.configure(state="normal", text="暂停")
        self.skip_btn.configure(state="normal")
        self.retry_btn.configure(state="disabled")
        self.worker.start()

    def toggle_pause(self) -> None:
        if not self.worker or not self.worker.is_alive():
            return
        if self.worker.pause_event.is_set():
            self.worker.pause_event.clear()
            self.pause_btn.configure(text="暂停")
            self.status_var.set("继续运行")
            self.log_line("继续")
        else:
            self.worker.pause_event.set()
            self.pause_btn.configure(text="继续")
            self.status_var.set("已暂停；不会保存或推进到下一票")
            self.log_line("暂停")

    def retry_current(self) -> None:
        if self.worker and self.worker.is_alive():
            self.worker.retry_event.set()
            self.retry_btn.configure(state="disabled")
            self.status_var.set("重新检测当前页面……")
            self.log_line("人工选择：重试当前")

    def skip_current(self) -> None:
        if self.worker and self.worker.is_alive():
            self.worker.skip_event.set()
            self.retry_btn.configure(state="disabled")
            self.log_line("人工选择：跳过当前")

    def recopy_current(self) -> None:
        number = self.current_var.get().strip()
        if number and number != "—":
            self.copy_to_clipboard(number)
            self.status_var.set(f"已重新复制：{number}")

    def stop(self) -> None:
        if self.worker and self.worker.is_alive():
            self.worker.stop_event.set()
            self.worker.retry_event.set()
            self.worker.skip_event.set()
            self.status_var.set("正在停止；不会再推进下一票")
            self.log_line("请求停止")

    def on_close(self) -> None:
        self.stop()
        self.root.after(200, self.root.destroy)

    def poll_ui_queue(self) -> None:
        try:
            while True:
                kind, data = self.ui_queue.get_nowait()
                if kind == "clipboard":
                    self.copy_to_clipboard(data.get("text", ""))
                elif kind == "current":
                    number = data["number"]
                    index = int(data["index"])
                    total = int(data["total"])
                    self.current_var.set(number)
                    self.progress_var.set(f"{index} / {total}")
                    self.progress["value"] = (index / total * 100) if total else 0
                    self.log_line(f"当前 {index}/{total}：{number}（已复制）")
                elif kind == "status":
                    msg = data.get("message", "")
                    self.status_var.set(msg)
                    self.log_line(msg)
                elif kind == "error":
                    msg = data.get("message", "")
                    self.status_var.set("⚠ " + msg)
                    self.retry_btn.configure(state="normal")
                    self.log_line("异常停止：" + msg)
                elif kind == "loaded":
                    total = int(data.get("total", 0))
                    self.progress_var.set(f"0 / {total}")
                    self.result_var.set(data.get("result_path", ""))
                    self.log_line(f"待处理 {total} 票；PDF目录：{data.get('output_dir', '')}")
                elif kind == "fatal":
                    msg = data.get("message", "")
                    self.status_var.set("运行失败：" + msg)
                    self.log_line("运行失败：" + msg)
                    self.reset_buttons()
                    messagebox.showerror(APP_TITLE, msg)
                elif kind == "finished":
                    msg = data.get("message", "")
                    self.status_var.set(msg)
                    self.log_line(msg)
                    self.reset_buttons()
        except queue.Empty:
            pass
        self.root.after(120, self.poll_ui_queue)

    def reset_buttons(self) -> None:
        self.start_btn.configure(state="normal")
        self.pause_btn.configure(state="disabled", text="暂停")
        self.retry_btn.configure(state="disabled")
        self.skip_btn.configure(state="disabled")


def main() -> int:
    parser = argparse.ArgumentParser(description="FedEx POD 半自动助手 - CDP打印版")
    parser.add_argument("excel", nargs="?", default="", help="可选：Excel 文件路径")
    parser.add_argument("--output", default="", help="可选：PDF 输出目录")
    parser.add_argument("--browser", choices=["edge", "chrome"], default="edge")
    parser.add_argument("--port", type=int, default=0)
    args = parser.parse_args()

    root = tk.Tk()
    app = App(root)
    if args.excel:
        app.excel_var.set(str(Path(args.excel).expanduser()))
    if args.output:
        app.output_var.set(str(Path(args.output).expanduser()))
    app.browser_var.set(args.browser.title())
    app.port_var.set(str(args.port or DEFAULT_PORTS[args.browser]))
    root.mainloop()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
